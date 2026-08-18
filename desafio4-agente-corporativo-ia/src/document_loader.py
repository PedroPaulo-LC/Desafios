from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import BinaryIO

import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".pptx", ".md", ".txt", ".csv", ".json", ".html", ".htm"}


def _read_bytes(file_obj: BinaryIO) -> bytes:
    file_obj.seek(0)
    data = file_obj.read()
    file_obj.seek(0)
    return data


def extract_text(file_obj: BinaryIO, filename: str) -> str:
    """Extrai texto de um arquivo corporativo suportado."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Formato não suportado: {extension}")

    data = _read_bytes(file_obj)

    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if extension == ".docx":
        document = Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs if p.text.strip())

    if extension in {".xlsx", ".xls"}:
        workbook = pd.ExcelFile(io.BytesIO(data))
        sections: list[str] = []
        for sheet in workbook.sheet_names:
            df = pd.read_excel(workbook, sheet_name=sheet)
            sections.append(f"Planilha: {sheet}\n{df.to_csv(index=False)}")
        return "\n\n".join(sections)

    if extension == ".pptx":
        presentation = Presentation(io.BytesIO(data))
        slides: list[str] = []
        for number, slide in enumerate(presentation.slides, start=1):
            parts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if parts:
                slides.append(f"Slide {number}:\n" + "\n".join(parts))
        return "\n\n".join(slides)

    text = data.decode("utf-8", errors="ignore")

    if extension == ".csv":
        rows = list(csv.reader(io.StringIO(text)))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)

    if extension == ".json":
        parsed = json.loads(text)
        return json.dumps(parsed, ensure_ascii=False, indent=2)

    if extension in {".html", ".htm"}:
        soup = BeautifulSoup(text, "html.parser")
        return soup.get_text("\n", strip=True)

    return text
